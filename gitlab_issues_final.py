"""
GitLab Issues Fetcher - FINAL OPTIMIZED VERSION
Features: Parallel processing, COPY inserts, logging, linked project names, proper column order
"""

import requests
import pandas as pd
import getpass
from typing import List, Dict, Any, Tuple
from datetime import datetime
import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from concurrent.futures import ThreadPoolExecutor, as_completed
from io import StringIO
import warnings
import logging
import sys
warnings.filterwarnings('ignore')

# Configure logging
log_filename = f'gitlab_fetch_{datetime.now().strftime("%Y%m%d_%H%M%S")}.log'
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler(log_filename),
        logging.StreamHandler(sys.stdout)
    ]
)
logger = logging.getLogger(__name__)


class GitLabIssuesFetcher:
    """Fetch all issues from GitLab projects with optimized parallel processing"""
    
    def __init__(self, gitlab_url: str, ikg_project_path: str, swat_project_path: str):
        self.gitlab_url = gitlab_url.rstrip('/')
        self.ikg_project_path = ikg_project_path
        self.swat_project_path = swat_project_path
        self.token = None
        self.headers = None
        self.project_id_to_name = {}  # Map project IDs to names
        logger.info("GitLabIssuesFetcher initialized")
        
    def authenticate(self):
        """Get authentication token from user"""
        self.token = getpass.getpass("Enter your GitLab Personal Access Token: ")
        self.headers = {
            'PRIVATE-TOKEN': self.token,
            'Content-Type': 'application/json'
        }
        logger.info("GitLab authentication configured")
        print("✓ GitLab authentication configured")
        
    def get_project_id(self, project_path: str) -> Tuple[str, str]:
        """Get project ID and name from project path"""
        logger.info(f"Fetching project info for: {project_path}")
        encoded_path = requests.utils.quote(project_path, safe='')
        url = f"{self.gitlab_url}/api/v4/projects/{encoded_path}"
        
        response = requests.get(url, headers=self.headers)
        response.raise_for_status()
        project_data = response.json()
        
        project_id = project_data['id']
        project_name = project_data['name']
        
        # Store mapping
        self.project_id_to_name[str(project_id)] = project_name
        
        logger.info(f"Found project: {project_name} (ID: {project_id})")
        print(f"✓ Found project: {project_name} (ID: {project_id})")
        return project_id, project_name
    
    def get_project_name_by_id(self, project_id: str) -> str:
        """Get project name by ID, fetch if not in cache"""
        project_id_str = str(project_id)
        
        if project_id_str in self.project_id_to_name:
            return self.project_id_to_name[project_id_str]
        
        # Fetch project info
        try:
            url = f"{self.gitlab_url}/api/v4/projects/{project_id}"
            response = requests.get(url, headers=self.headers, timeout=5)
            response.raise_for_status()
            project_data = response.json()
            project_name = project_data['name']
            self.project_id_to_name[project_id_str] = project_name
            logger.debug(f"Fetched project name for ID {project_id}: {project_name}")
            return project_name
        except Exception as e:
            logger.warning(f"Could not fetch project name for ID {project_id}: {e}")
            return f"Project_{project_id}"
    
    def fetch_issue_links(self, project_id: str, issue_iid: int) -> List[Dict[str, Any]]:
        """Fetch all links for a specific issue"""
        url = f"{self.gitlab_url}/api/v4/projects/{project_id}/issues/{issue_iid}/links"
        
        try:
            response = requests.get(url, headers=self.headers, timeout=10)
            response.raise_for_status()
            return response.json()
        except Exception as e:
            logger.debug(f"No links for issue {issue_iid}: {e}")
            return []
    
    def fetch_links_batch(self, project_id: str, issue_iids: List[int], max_workers: int = 10) -> Dict[int, List[Dict]]:
        """Fetch links for multiple issues in parallel"""
        logger.info(f"Fetching links for {len(issue_iids)} issues using {max_workers} threads")
        links_map = {}
        
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_iid = {
                executor.submit(self.fetch_issue_links, project_id, iid): iid 
                for iid in issue_iids
            }
            
            completed = 0
            for future in as_completed(future_to_iid):
                iid = future_to_iid[future]
                try:
                    links_map[iid] = future.result()
                    completed += 1
                    if completed % 100 == 0:
                        logger.info(f"Fetched links for {completed}/{len(issue_iids)} issues")
                except Exception as e:
                    logger.warning(f"Error fetching links for issue {iid}: {e}")
                    links_map[iid] = []
        
        logger.info(f"Completed fetching links for all {len(issue_iids)} issues")
        return links_map
    
    def fetch_all_issues(self, project_id: str, project_name: str) -> List[Dict[str, Any]]:
        """Fetch all issues with optimized parallel link fetching"""
        all_issues = []
        page = 1
        per_page = 100
        
        url = f"{self.gitlab_url}/api/v4/projects/{project_id}/issues"
        
        logger.info(f"Starting to fetch issues from {project_name}")
        print(f"\nFetching issues from {project_name}...")
        
        # Step 1: Fetch all issues
        while True:
            params = {
                'per_page': per_page,
                'page': page,
                'state': 'all',
                'scope': 'all',
                'with_labels_details': True
            }
            
            try:
                response = requests.get(url, headers=self.headers, params=params)
                response.raise_for_status()
                issues = response.json()
                
                if not issues:
                    break
                
                all_issues.extend(issues)
                logger.info(f"Page {page}: fetched {len(issues)} issues (Total: {len(all_issues)})")
                print(f"  Page {page}: {len(issues)} issues (Total: {len(all_issues)})")
                
                if len(issues) < per_page:
                    break
                
                page += 1
            except Exception as e:
                logger.error(f"Error fetching issues page {page}: {e}")
                raise
        
        logger.info(f"Fetched {len(all_issues)} issues from {project_name}")
        print(f"✓ Fetched {len(all_issues)} issues")
        
        # Step 2: Fetch all links in parallel
        if all_issues:
            print(f"Fetching links in parallel (using 20 threads)...")
            issue_iids = [issue['iid'] for issue in all_issues]
            links_map = self.fetch_links_batch(project_id, issue_iids, max_workers=20)
            
            # Attach links to issues
            for issue in all_issues:
                issue['_links_data'] = links_map.get(issue['iid'], [])
            
            logger.info("Links attached to all issues")
            print(f"✓ Links fetched")
        
        return all_issues
    
    def extract_issue_data(self, issues: List[Dict[str, Any]], project_identifier: str) -> pd.DataFrame:
        """Extract comprehensive issue data with proper column order"""
        logger.info(f"Extracting data from {len(issues)} issues for {project_identifier}")
        
        # Debug: Log structure of first issue if available
        if issues and len(issues) > 0:
            first_issue = issues[0]
            logger.debug(f"Sample issue keys: {list(first_issue.keys())}")
            if first_issue.get('iteration'):
                logger.info(f"Sample iteration data: {first_issue.get('iteration')}")
            else:
                logger.debug("No iteration field in sample issue")
        
        extracted_data = []
        current_datetime = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        
        for issue in issues:
            # Extract assignees
            assignees = [a.get('name', '') for a in issue.get('assignees', [])]
            assignee_str = ', '.join(assignees) if assignees else None
            
            # Extract labels (handle both string and dict formats)
            labels = issue.get('labels', [])
            if labels:
                if isinstance(labels[0], dict):
                    label_str = ', '.join([label.get('name', '') for label in labels])
                else:
                    label_str = ', '.join(labels)
            else:
                label_str = None
            
            # Extract iteration - try multiple possible field structures
            iteration = None
            if issue.get('iteration'):
                iteration_data = issue.get('iteration')
                # Try different possible fields for iteration name/title
                if isinstance(iteration_data, dict):
                    iteration = (iteration_data.get('title') or 
                               iteration_data.get('name') or 
                               iteration_data.get('web_url', '').split('/')[-1] if iteration_data.get('web_url') else None)
                elif isinstance(iteration_data, str):
                    iteration = iteration_data
                
                if iteration:
                    logger.debug(f"Issue {issue.get('iid')} iteration: {iteration}")
            
            epic = issue.get('epic', {}).get('title', '') if issue.get('epic') else None
            epic_iid = issue.get('epic', {}).get('iid', '') if issue.get('epic') else None
            
            participants = [p.get('name', '') for p in issue.get('participants', [])]
            participants_str = ', '.join(participants) if participants else None
            
            milestone = issue.get('milestone', {}).get('title', '') if issue.get('milestone') else None
            
            time_estimate = issue.get('time_stats', {}).get('time_estimate')
            time_spent = issue.get('time_stats', {}).get('total_time_spent')
            
            task_completion = None
            if issue.get('task_completion_status'):
                completed = issue['task_completion_status'].get('completed_count', 0)
                total = issue['task_completion_status'].get('count', 0)
                task_completion = f"{completed}/{total}"
            
            # Consolidate link data and get linked project names
            links_data = issue.get('_links_data', [])
            
            link_ids = []
            link_issue_ids = []
            link_issue_iids = []
            link_types = []
            link_urls = []
            link_issue_titles = []
            linked_project_ids = []
            linked_project_names = []
            
            for link in links_data:
                if link.get('id'):
                    link_ids.append(str(link.get('id', '')))
                if link.get('issue_link_id'):
                    link_issue_ids.append(str(link.get('issue_link_id', '')))
                if link.get('iid'):
                    link_issue_iids.append(str(link.get('iid', '')))
                if link.get('link_type'):
                    link_types.append(str(link.get('link_type', '')))
                if link.get('web_url'):
                    link_urls.append(str(link.get('web_url', '')))
                if link.get('title'):
                    link_issue_titles.append(str(link.get('title', '')))
                if link.get('project_id'):
                    proj_id = str(link.get('project_id', ''))
                    linked_project_ids.append(proj_id)
                    # Get project name
                    proj_name = self.get_project_name_by_id(proj_id)
                    linked_project_names.append(proj_name)
            
            # Column order: project, issue_id, issue_iid, title, labels, description, ...
            issue_data = {
                'project': project_identifier,
                'issue_id': issue.get('id'),
                'issue_iid': issue.get('iid'),
                'title': issue.get('title'),
                'labels': label_str,  # Labels right after title
                'description': issue.get('description'),
                'state': issue.get('state'),
                'web_url': issue.get('web_url', ''),
                
                # Linkage columns with linked_project_name
                'link_id': ', '.join(link_ids) if link_ids else None,
                'link_issue_id': ', '.join(link_issue_ids) if link_issue_ids else None,
                'link_issue_iid': ', '.join(link_issue_iids) if link_issue_iids else None,
                'link_type': ', '.join(link_types) if link_types else None,
                'link_url': ', '.join(link_urls) if link_urls else None,
                'link_issue_title': ', '.join(link_issue_titles) if link_issue_titles else None,
                'linked_project_id': ', '.join(linked_project_ids) if linked_project_ids else None,
                'linked_project_name': ', '.join(linked_project_names) if linked_project_names else None,  # NEW!
                
                # Assignment and ownership
                'author': issue.get('author', {}).get('name'),
                'author_username': issue.get('author', {}).get('username'),
                'created_by_id': issue.get('author', {}).get('id'),
                'assignee': assignee_str,
                'assignee_ids': ', '.join([str(a.get('id', '')) for a in issue.get('assignees', [])]),
                
                # Dates
                'issue_created_date': issue.get('created_at'),
                'created_at': issue.get('created_at'),
                'updated_at': issue.get('updated_at'),
                'closed_at': issue.get('closed_at'),
                'due_date': issue.get('due_date'),
                'start_date': issue.get('start_date'),
                'current_date_time': current_datetime,
                
                # Organization
                'milestone': milestone,
                'iteration': iteration,
                'epic': epic,
                'epic_iid': epic_iid,
                'weight': issue.get('weight'),
                
                # Relationships
                'parent_iid': None,
                'has_tasks': issue.get('has_tasks'),
                'task_completion_status': task_completion,
                
                # Engagement
                'participants': participants_str,
                'upvotes': issue.get('upvotes'),
                'downvotes': issue.get('downvotes'),
                'user_notes_count': issue.get('user_notes_count'),
                'merge_requests_count': issue.get('merge_requests_count'),
                
                # Time tracking
                'time_estimate_hours': time_estimate / 3600 if time_estimate else None,
                'time_spent_hours': time_spent / 3600 if time_spent else None,
                
                # Metadata
                'confidential': issue.get('confidential'),
                'discussion_locked': issue.get('discussion_locked'),
                'issue_type': issue.get('issue_type'),
                'severity': issue.get('severity'),
                'health_status': issue.get('health_status'),
            }
            
            extracted_data.append(issue_data)
        
        df = pd.DataFrame(extracted_data)
        
        # Log iteration statistics
        iterations_found = df['iteration'].notna().sum()
        logger.info(f"Extracted {len(df)} rows with {len(df.columns)} columns")
        logger.info(f"Issues with iteration data: {iterations_found}/{len(df)}")
        
        if iterations_found > 0:
            unique_iterations = df[df['iteration'].notna()]['iteration'].unique()
            logger.info(f"Unique iterations found: {list(unique_iterations)[:10]}")  # Log first 10
        else:
            logger.warning("No iteration data found in any issues - this may be expected if issues don't have iterations assigned")
        
        return df


class GreenplumLoader:
    """Load data into Greenplum with optimized bulk operations"""
    
    def __init__(self, host: str, port: int, database: str, user: str, schema: str):
        self.host = host
        self.port = port
        self.database = database
        self.user = user
        self.schema = schema
        self.password = None
        self.conn = None
        logger.info(f"GreenplumLoader initialized for {database}.{schema}")
        
    def get_password(self):
        """Get database password from user"""
        self.password = getpass.getpass("Enter Greenplum database password: ")
        logger.info("Database password configured")
        print("✓ Database password configured")
        
    def connect(self):
        """Establish connection to Greenplum"""
        try:
            logger.info(f"Connecting to Greenplum: {self.host}:{self.port}/{self.database}")
            self.conn = psycopg2.connect(
                host=self.host,
                port=self.port,
                database=self.database,
                user=self.user,
                password=self.password
            )
            self.conn.autocommit = False
            logger.info("Connected to Greenplum database")
            print(f"✓ Connected to Greenplum database: {self.database}")
        except Exception as e:
            logger.error(f"Error connecting to database: {e}")
            print(f"✗ Error connecting to database: {e}")
            raise
    
    def create_table(self, table_name: str, df: pd.DataFrame, drop_if_exists: bool = True, if_not_exists: bool = False):
        """
        Create table in Greenplum
        - For main tables: drop_if_exists=True (drop and recreate every time)
        - For archive tables: if_not_exists=True (create only if not exists)
        """
        cursor = self.conn.cursor()
        
        try:
            # Drop table if requested (for main tables)
            if drop_if_exists:
                drop_query = f"DROP TABLE IF EXISTS {self.schema}.{table_name}"
                cursor.execute(drop_query)
                logger.info(f"Dropped table {self.schema}.{table_name} (if existed)")
                print(f"  Dropped {self.schema}.{table_name} (if existed)")
            
            # Generate CREATE TABLE statement
            columns = []
            for col, dtype in df.dtypes.items():
                if dtype == 'object':
                    col_type = 'TEXT'
                elif dtype == 'int64':
                    col_type = 'BIGINT'
                elif dtype == 'float64':
                    col_type = 'DOUBLE PRECISION'
                elif dtype == 'bool':
                    col_type = 'BOOLEAN'
                elif dtype == 'datetime64[ns]':
                    col_type = 'TIMESTAMP'
                else:
                    col_type = 'TEXT'
                
                columns.append(f"{col} {col_type}")
            
            columns_str = ', '.join(columns)
            
            # Use IF NOT EXISTS for archive tables
            if if_not_exists:
                create_query = f"""
                    CREATE TABLE IF NOT EXISTS {self.schema}.{table_name} (
                        {columns_str}
                    )
                    DISTRIBUTED RANDOMLY
                """
            else:
                create_query = f"""
                    CREATE TABLE {self.schema}.{table_name} (
                        {columns_str}
                    )
                    DISTRIBUTED RANDOMLY
                """
            
            cursor.execute(create_query)
            self.conn.commit()
            logger.info(f"Created table {self.schema}.{table_name}")
            print(f"  ✓ Created {self.schema}.{table_name}")
            
        except Exception as e:
            self.conn.rollback()
            if if_not_exists and 'already exists' in str(e).lower():
                logger.info(f"Table {self.schema}.{table_name} already exists")
                print(f"  ✓ Table {self.schema}.{table_name} already exists")
                self.conn.commit()
            else:
                logger.error(f"Error creating table {table_name}: {e}")
                print(f"  ✗ Error: {e}")
                raise
        finally:
            cursor.close()
    
    def insert_data(self, table_name: str, df: pd.DataFrame):
        """Insert data using ultra-fast COPY command"""
        cursor = self.conn.cursor()
        
        try:
            logger.info(f"Inserting {len(df)} rows into {self.schema}.{table_name} using COPY")
            
            # Create CSV buffer in memory
            buffer = StringIO()
            df.to_csv(buffer, index=False, header=False, sep='\t', na_rep='\\N')
            buffer.seek(0)
            
            # Use COPY command - fastest way to load data
            columns_str = ', '.join(df.columns)
            copy_sql = f"""
                COPY {self.schema}.{table_name} ({columns_str})
                FROM STDIN WITH (FORMAT CSV, DELIMITER E'\\t', NULL '\\N')
            """
            
            cursor.copy_expert(copy_sql, buffer)
            self.conn.commit()
            
            logger.info(f"Inserted {len(df)} rows into {self.schema}.{table_name}")
            print(f"  ✓ Inserted {len(df)} rows into {self.schema}.{table_name}")
            
        except Exception as e:
            self.conn.rollback()
            logger.error(f"Error inserting data into {table_name}: {e}")
            print(f"  ✗ Error: {e}")
            raise
        finally:
            cursor.close()
    
    def close(self):
        """Close database connection"""
        if self.conn:
            self.conn.close()
            logger.info("Database connection closed")
            print("✓ Database connection closed")


def create_excel_with_sheets(ikg_df: pd.DataFrame, swat_df: pd.DataFrame, filename: str):
    """Create Excel file with two worksheets"""
    logger.info(f"Creating Excel file: {filename}")
    print("\nCreating Excel file...")
    
    wb = Workbook(write_only=False)
    wb.remove(wb.active)
    
    # Create IKG sheet
    logger.info("Creating IKG worksheet")
    ikg_sheet = wb.create_sheet('IKG')
    for r_idx, r in enumerate(dataframe_to_rows(ikg_df, index=False, header=True)):
        ikg_sheet.append(r)
        if r_idx == 0:
            for cell in ikg_sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
                cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Create SWAT sheet
    logger.info("Creating SWAT worksheet")
    swat_sheet = wb.create_sheet('SWAT')
    for r_idx, r in enumerate(dataframe_to_rows(swat_df, index=False, header=True)):
        swat_sheet.append(r)
        if r_idx == 0:
            for cell in swat_sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
                cell.alignment = Alignment(horizontal='center', vertical='center')
    
    wb.save(filename)
    logger.info(f"Excel file saved: {filename}")
    print(f"✓ Excel file saved: {filename}")


def main():
    """Main execution function"""
    
    print("=" * 80)
    print("GitLab Issues Fetcher - FINAL OPTIMIZED VERSION")
    print("Features: Parallel processing, COPY inserts, logging, linked project names")
    print("=" * 80)
    
    logger.info("="*80)
    logger.info("Starting GitLab Issues Fetch Process")
    logger.info("="*80)
    
    # Configuration
    GITLAB_URL = 'https://devcloud.ubs.net'
    IKG_PROJECT_PATH = 'ubs/gwma/smart-technology-and-analytics/staat-data-science/staat-ds-insights-cl/commons/staat-ds-insights-home'
    SWAT_PROJECT_PATH = 'ubs/gwma/smart-technology-and-analytics/staat-data-science/staat-ds-insights-cl/commons/staat-ds-insights-cl-home'
    
    GREENPLUM_HOST = 'greenplum-rdsp.zur.swissbank.com'
    GREENPLUM_PORT = 5432
    GREENPLUM_DB = 'gprdsp'
    GREENPLUM_USER = 'ds_rdsp_dev'
    GREENPLUM_SCHEMA = 'sandbox_prj_smart_insights'
    
    OUTPUT_TABLE1 = 'ikg_issue_details'
    OUTPUT_TABLE2 = 'swat_issue_details'
    ARCHIVE_TABLE1 = 'ikg_issue_details_archive'
    ARCHIVE_TABLE2 = 'swat_issue_details_archive'
    
    logger.info(f"Configuration: GitLab={GITLAB_URL}, Greenplum={GREENPLUM_HOST}:{GREENPLUM_PORT}/{GREENPLUM_DB}")
    logger.info(f"Log file: {log_filename}")
    print(f"Log file: {log_filename}")
    
    start_time = datetime.now()
    logger.info(f"Process started at {start_time}")
    
    # Step 1: Authenticate
    fetcher = GitLabIssuesFetcher(GITLAB_URL, IKG_PROJECT_PATH, SWAT_PROJECT_PATH)
    fetcher.authenticate()
    
    # Step 2: Fetch IKG issues
    logger.info("Starting IKG project fetch")
    ikg_project_id, ikg_project_name = fetcher.get_project_id(IKG_PROJECT_PATH)
    ikg_issues = fetcher.fetch_all_issues(ikg_project_id, ikg_project_name)
    ikg_df = fetcher.extract_issue_data(ikg_issues, 'staat-ds-insights-home')
    logger.info(f"IKG extraction complete: {len(ikg_df)} rows")
    print(f"✓ IKG: {len(ikg_df)} issues extracted")
    
    # Step 3: Fetch SWAT issues
    logger.info("Starting SWAT project fetch")
    swat_project_id, swat_project_name = fetcher.get_project_id(SWAT_PROJECT_PATH)
    swat_issues = fetcher.fetch_all_issues(swat_project_id, swat_project_name)
    swat_df = fetcher.extract_issue_data(swat_issues, 'staat-ds-insights-cl-home')
    logger.info(f"SWAT extraction complete: {len(swat_df)} rows")
    print(f"✓ SWAT: {len(swat_df)} issues extracted")
    
    # Step 4: Create Excel file
    excel_filename = 'gitlab_issues_ikg_swat.xlsx'
    create_excel_with_sheets(ikg_df, swat_df, excel_filename)
    
    # Step 5: Database operations
    print("\n" + "=" * 80)
    print("GREENPLUM DATABASE OPERATIONS")
    print("=" * 80)
    logger.info("Starting Greenplum database operations")
    
    loader = GreenplumLoader(GREENPLUM_HOST, GREENPLUM_PORT, GREENPLUM_DB, 
                            GREENPLUM_USER, GREENPLUM_SCHEMA)
    loader.get_password()
    loader.connect()
    
    # Step 6: Create main tables (DROP and CREATE every time)
    print("\nCreating main tables (drop and create)...")
    logger.info("Creating main tables (drop and create)")
    loader.create_table(OUTPUT_TABLE1, ikg_df, drop_if_exists=True, if_not_exists=False)
    loader.create_table(OUTPUT_TABLE2, swat_df, drop_if_exists=True, if_not_exists=False)
    
    # Step 7: Create archive tables (CREATE IF NOT EXISTS)
    print("\nCreating archive tables (if not exists)...")
    logger.info("Creating archive tables (if not exists)")
    loader.create_table(ARCHIVE_TABLE1, ikg_df, drop_if_exists=False, if_not_exists=True)
    loader.create_table(ARCHIVE_TABLE2, swat_df, drop_if_exists=False, if_not_exists=True)
    
    # Step 8: Insert data into main tables (using COPY)
    print("\nInserting data into main tables (COPY command)...")
    logger.info("Inserting data into main tables")
    loader.insert_data(OUTPUT_TABLE1, ikg_df)
    loader.insert_data(OUTPUT_TABLE2, swat_df)
    
    # Step 9: Insert data into archive tables (using COPY)
    print("\nInserting data into archive tables (COPY command)...")
    logger.info("Inserting data into archive tables")
    loader.insert_data(ARCHIVE_TABLE1, ikg_df)
    loader.insert_data(ARCHIVE_TABLE2, swat_df)
    
    loader.close()
    
    # Final summary
    end_time = datetime.now()
    duration = (end_time - start_time).total_seconds()
    
    print("\n" + "=" * 80)
    print("SUMMARY")
    print("=" * 80)
    print(f"IKG Issues: {len(ikg_df)} rows")
    print(f"SWAT Issues: {len(swat_df)} rows")
    print(f"Excel file: {excel_filename}")
    print(f"\nGreenplum Tables:")
    print(f"  Main Tables (drop & create):")
    print(f"    - {GREENPLUM_SCHEMA}.{OUTPUT_TABLE1}")
    print(f"    - {GREENPLUM_SCHEMA}.{OUTPUT_TABLE2}")
    print(f"  Archive Tables (if not exists):")
    print(f"    - {GREENPLUM_SCHEMA}.{ARCHIVE_TABLE1}")
    print(f"    - {GREENPLUM_SCHEMA}.{ARCHIVE_TABLE2}")
    print(f"\n⏱️  Total execution time: {duration:.2f} seconds ({duration/60:.2f} minutes)")
    print(f"📝 Log file: {log_filename}")
    print("\n✓ Process completed successfully!")
    
    logger.info("="*80)
    logger.info(f"Process completed successfully in {duration:.2f} seconds")
    logger.info(f"IKG: {len(ikg_df)} rows, SWAT: {len(swat_df)} rows")
    logger.info("="*80)


if __name__ == "__main__":
    main()
